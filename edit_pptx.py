import copy
from pptx import Presentation
from pptx.util import Emu

SRC = "/home/ubuntu/Workspace/Amodal_Completion_Presentation.pptx"
OUT = "/home/ubuntu/Workspace/Amodal_Completion_Presentation.new.pptx"

CAT_BUCKET_IMG = "/home/ubuntu/Workspace/amodal_completion/output/cat-7875506_640_cat/_flux_cutout_cat/final/comparison.png"
CAT_BUCKET_W, CAT_BUCKET_H = 3200, 388

prs = Presentation(SRC)


def slide_by_title(prs, title_substr):
    """Match only the slide's TITLE shape (the first text-bearing shape),
    not any text anywhere on the slide — avoids matching e.g. the Agenda
    slide's bullet list, which repeats every other slide's title text."""
    for i, s in enumerate(prs.slides):
        for shp in s.shapes:
            if shp.has_text_frame and shp.text_frame.text.strip():
                if title_substr.lower() in shp.text_frame.text.lower():
                    return i, s
                break  # first non-empty text shape = title; stop here either way
    raise ValueError(f"no slide found with title containing {title_substr!r}")


def shape_by_id(slide, shape_id):
    for shp in slide.shapes:
        if shp.shape_id == shape_id:
            return shp
    raise ValueError(f"no shape id={shape_id} on this slide")


def set_text_keep_format(shape, new_text):
    """Overwrite a text frame's text while keeping the first run's formatting."""
    tf = shape.text_frame
    first_para = tf.paragraphs[0]
    lines = new_text.split("\n")
    # Reuse formatting off the very first run in the whole text frame.
    template_run = None
    for para in tf.paragraphs:
        if para.runs:
            template_run = para.runs[0]
            break
    # Clear all paragraphs except the first, then clear first paragraph's runs.
    for para in list(tf.paragraphs[1:]):
        para._p.getparent().remove(para._p)
    for r in list(first_para.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        para = first_para if i == 0 else tf.add_paragraph()
        if i > 0:
            # copy paragraph-level formatting from first_para
            para._p.set('marL', first_para._p.get('marL')) if first_para._p.get('marL') else None
        run = para.add_run()
        run.text = line
        if template_run is not None:
            run.font.size = template_run.font.size
            run.font.bold = template_run.font.bold
            run.font.name = template_run.font.name
            if template_run.font.color and template_run.font.color.type is not None:
                run.font.color.rgb = template_run.font.color.rgb


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def duplicate_slide_structure(prs, src_slide, skip_shape_ids=()):
    """Append a new slide using the same layout, copy all shapes from
    src_slide except those in skip_shape_ids (deepcopy — safe for shapes
    that don't reference external relationship parts, i.e. no pictures)."""
    dest = prs.slides.add_slide(src_slide.slide_layout)
    for shp in list(dest.shapes):
        remove_shape(shp)
    for shp in src_slide.shapes:
        if shp.shape_id in skip_shape_ids:
            continue
        new_el = copy.deepcopy(shp._element)
        dest.shapes._spTree.append(new_el)
    return dest


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])


def delete_slide_by_index(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


# ── 1. Slide 1 — remove guide name, drop "Multi-Agent" from subtitle ────────
i1, s1 = slide_by_title(prs, "Amodal Completion")
remove_shape(shape_by_id(s1, 8))  # "Guide: Deepa Sreedhar"
subtitle = shape_by_id(s1, 5)
set_text_keep_format(subtitle,
    "An Automated Pipeline for Occlusion-Aware\nImage Reconstruction Using Segmentation and Diffusion Models")
print("Slide 1 done")

# ── 2. Slide 4 — drop "optional text hint" wording ──────────────────────────
i4, s4 = slide_by_title(prs, "Objectives & Scope")
autodetect = shape_by_id(s4, 9)
set_text_keep_format(autodetect, "Identify occluded subject & occluder automatically.")
print("Slide 4 done")

# ── 3. Slide 6 — rewrite architecture, no LangGraph / Jiang Ao / LISA / Hunyuan3D
i6, s6 = slide_by_title(prs, "System Architecture")
left = shape_by_id(s6, 6)
right = shape_by_id(s6, 7)
set_text_keep_format(left,
    "›  Single Python pipeline (pipeline.py) — sequential stages\n"
    "›  Agent 1: GPT-5 vision names subject & occluder\n"
    "›  SAM3 text-prompted segmentation (primary) + InstaFormer IoU-matched occlusion order\n"
    "›  Priority-based mask fusion + area-ratio oversized-mask trim")
set_text_keep_format(right,
    "›  Dilated-occluder hidden-region formula: dilate(occluder) \\ visible\n"
    "›  FLUX.1-Fill-dev inpainting, GPT-vision reviewer + retry loop\n"
    "›  SAM3 post-segmentation refinement of Flux output\n"
    "›  Final RGBA composite")
print("Slide 6 done")

# ── 4. Slide 8 — drop "(Jiang Ao)" attribution ───────────────────────────────
i8, s8 = slide_by_title(prs, "Mask Fusion & Hidden-Region")
formula_title = shape_by_id(s8, 9)
set_text_keep_format(formula_title, "Hidden-Region Formula")
print("Slide 8 done")

# ── 5. Slide 12 (Fox) — relabel as Failure 2 ─────────────────────────────────
i12, s12 = slide_by_title(prs, "Result: Fox")
fox_title = shape_by_id(s12, 4)
fox_caption = shape_by_id(s12, 7)
set_text_keep_format(fox_title, "Failure 2: Fox")
set_text_keep_format(fox_caption,
    "Snowbank occluder removed, but the front leg is only partially shown — full leg anatomy not generated.")
print("Slide 12 relabeled as Failure 2")

# ── 6. New slide "Failure 1: Cat" — duplicate slide 12's structure, own image
new_cat_fail = duplicate_slide_structure(prs, s12, skip_shape_ids={6})
new_title = shape_by_id(new_cat_fail, 4)
new_caption = shape_by_id(new_cat_fail, 7)
set_text_keep_format(new_title, "Failure 1: Cat")
set_text_keep_format(new_caption,
    "Bucket occluder removed, but the cat's back/hindquarters are only partially shown — not fully generated.")
# place the picture at the same box the template used
template_pic = shape_by_id(s12, 6)
left_pos, top_pos, w_pos = template_pic.left, template_pic.top, template_pic.width
h_pos = int(w_pos * CAT_BUCKET_H / CAT_BUCKET_W)
new_cat_fail.shapes.add_picture(CAT_BUCKET_IMG, left_pos, top_pos, width=w_pos, height=h_pos)
# move it to right after the current index of the Cat(bowl) slide, i.e. right before Fox
i_cat_bowl, _ = slide_by_title(prs, "Result: Cat (Eating Food)")
new_index_pos = i_cat_bowl + 1
current_new_slide_index = len(prs.slides.__iter__.__self__._sldIdLst) - 1  # last position
move_slide(prs, len(prs.slides) - 1, new_index_pos)
print("New Failure 1: Cat slide inserted")

# ── 8. New "Implementation Details" slide, right after "Implementation Highlights"
i_impl_hl, s_impl_hl = slide_by_title(prs, "Implementation Highlights")
# Use "Future Work" as the simple single-column template (title + 1 bullet box + page#)
_, s_future_tmpl = slide_by_title(prs, "Future Work")
new_impl = duplicate_slide_structure(prs, s_future_tmpl)
impl_title = shape_by_id(new_impl, 4)
impl_body = shape_by_id(new_impl, 5)
set_text_keep_format(impl_title, "Implementation Details")
set_text_keep_format(impl_body,
    "›  Implemented in Python; dependencies managed with uv (pyproject.toml + uv.lock)\n"
    "›  Tested on openly available stock-photo images (Pexels / Pixabay) — not a fixed academic benchmark\n"
    "›  No fine-tuning or training — every model used as a pretrained, open-source checkpoint\n"
    "›  Evaluated manually via human + GPT-vision review; no automated ground-truth metric\n"
    "›  All stages run on a single RTX A6000 48GB GPU")
move_slide(prs, len(prs.slides) - 1, i_impl_hl + 1)
print("New Implementation Details slide inserted")

# ── 9. New "Real-World Applications" slide, right after "Conclusion & Contributions"
i_concl, s_concl = slide_by_title(prs, "Conclusion & Contributions")
new_apps = duplicate_slide_structure(prs, s_future_tmpl)
apps_title = shape_by_id(new_apps, 4)
apps_body = shape_by_id(new_apps, 5)
set_text_keep_format(apps_title, "Real-World Applications")
set_text_keep_format(apps_body,
    "›  Autonomous vehicles — reconstructing occluded pedestrians/vehicles for safer perception\n"
    "›  Medical robotics — inferring occluded anatomy or instruments during surgical guidance\n"
    "›  Photo & heritage restoration — completing damaged or obstructed archival images\n"
    "›  E-commerce & product photography — removing hands/packaging from product shots")
move_slide(prs, len(prs.slides) - 1, i_concl + 1)
print("New Real-World Applications slide inserted")

# ── 10. Last slide — remove "Questions?" ─────────────────────────────────────
i_last, s_last = slide_by_title(prs, "Thank You")
remove_shape(shape_by_id(s_last, 4))
print("Removed 'Questions?' from last slide")

# ── 7 (moved here). Delete "Results Summary Across Categories" slide ────────
# Done LAST, right before save — deleting a slide's part before later
# add_slide() calls can leave its old part-name slot in a state where a
# subsequently-added slide's part collides with it, producing a duplicate
# zip entry. Doing every addition first, deletion last, avoids the collision.
i_summary, _ = slide_by_title(prs, "Results Summary Across Categories")
delete_slide_by_index(prs, i_summary)
print(f"Deleted Results Summary slide (was index {i_summary})")

# ── 11. Renumber every page-number footer (the last TEXT-bearing shape on
#        every slide except the first [title] and last [thank you]) to match
#        final slide order. Search backward past any non-text shapes (e.g.
#        a picture appended after the footer during slide construction).
n = len(prs.slides)
for idx, slide in enumerate(prs.slides):
    if idx == 0 or idx == n - 1:
        continue
    shapes = list(slide.shapes)
    footer = None
    for shp in reversed(shapes):
        if shp.has_text_frame and shp.text_frame.text.strip():
            footer = shp
            break
    if footer is not None:
        set_text_keep_format(footer, str(idx + 1))
    else:
        print(f"  [WARN] no footer text shape found on slide {idx + 1}")
print("Renumbered page footers")

prs.save(OUT)
print("Saved to", OUT)
