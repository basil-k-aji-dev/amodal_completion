import copy
from pptx import Presentation
from pptx.util import Emu

PPTX = "/home/ubuntu/Workspace/Amodal_Completion_Presentation.pptx"
ICON_DIR = "/home/ubuntu/Workspace/amodal_completion/icons"
DIAGRAM = "/home/ubuntu/Workspace/amodal_completion/architecture_diagram.png"

prs = Presentation(PPTX)


def slide_by_title(prs, title_substr):
    for i, s in enumerate(prs.slides):
        for shp in s.shapes:
            if shp.has_text_frame and shp.text_frame.text.strip():
                if title_substr.lower() in shp.text_frame.text.lower():
                    return i, s
                break
    raise ValueError(f"no slide found with title containing {title_substr!r}")


def shape_by_id(slide, shape_id):
    for shp in slide.shapes:
        if shp.shape_id == shape_id:
            return shp
    raise ValueError(f"no shape id={shape_id} on this slide")


def set_text_keep_format(shape, new_text):
    tf = shape.text_frame
    first_para = tf.paragraphs[0]
    lines = new_text.split("\n")
    template_run = None
    for para in tf.paragraphs:
        if para.runs:
            template_run = para.runs[0]
            break
    for para in list(tf.paragraphs[1:]):
        para._p.getparent().remove(para._p)
    for r in list(first_para.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        para = first_para if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = line
        if template_run is not None:
            run.font.size = template_run.font.size
            run.font.bold = template_run.font.bold
            run.font.name = template_run.font.name
            if template_run.font.color and template_run.font.color.type is not None:
                run.font.color.rgb = template_run.font.color.rgb


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def duplicate_slide_structure(prs, src_slide, skip_shape_ids=()):
    dest = prs.slides.add_slide(src_slide.slide_layout)
    for shp in list(dest.shapes):
        remove_shape(shp)
    for shp in src_slide.shapes:
        if shp.shape_id in skip_shape_ids:
            continue
        new_el = copy.deepcopy(shp._element)
        dest.shapes._spTree.append(new_el)
    return dest


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])


def delete_slide_by_index(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


# ── 1. Replace slide 6's architecture diagram with the fixed high-res version
i6, s6 = slide_by_title(prs, "System Architecture")
old_pic = None
for shp in s6.shapes:
    if shp.shape_type == 13:
        old_pic = shp
        break
left, top, width, height = old_pic.left, old_pic.top, old_pic.width, old_pic.height
remove_shape(old_pic)
s6.shapes.add_picture(DIAGRAM, left, top, width=width, height=height)
print("Slide 6 diagram replaced")

# ── 2. Rebuild "Real-World Applications" as a 4-card icon grid ──────────────
i_apps, s_apps = slide_by_title(prs, "Real-World Applications")
_, s_cards_tmpl = slide_by_title(prs, "Objectives & Scope")

new_apps = duplicate_slide_structure(prs, s_cards_tmpl)
set_text_keep_format(shape_by_id(new_apps, 4), "Real-World Applications")

cards = [
    # (number_shape_id, heading_id, body_id, icon_file, heading_text, body_text)
    (7,  8,  9,  "car.png",     "Autonomous Vehicles",
     "Reconstruct occluded pedestrians and vehicles for safer perception."),
    (12, 13, 14, "medical.png", "Medical Robotics",
     "Infer occluded anatomy or instruments during surgical guidance."),
    (17, 18, 19, "photo.png",   "Heritage Restoration",
     "Complete damaged or obstructed archival photographs."),
    (22, 23, 24, "cart.png",    "E-Commerce Photography",
     "Remove hands or packaging from product shots."),
]

for number_id, heading_id, body_id, icon_file, heading_text, body_text in cards:
    number_shape = shape_by_id(new_apps, number_id)
    icon_size = Emu(731520)  # square, matches the old number textbox height
    icon_left = Emu(int(number_shape.left + (number_shape.width - icon_size) / 2))
    icon_top = number_shape.top
    remove_shape(number_shape)
    new_apps.shapes.add_picture(f"{ICON_DIR}/{icon_file}", icon_left, icon_top,
                                 width=icon_size, height=icon_size)
    set_text_keep_format(shape_by_id(new_apps, heading_id), heading_text)
    set_text_keep_format(shape_by_id(new_apps, body_id), body_text)

# remove the old plain-bullet Real-World Applications slide, put the new one in its place
delete_slide_by_index(prs, i_apps)
move_slide(prs, len(prs.slides) - 1, i_apps)
print("Real-World Applications rebuilt as icon-card grid")

# ── 3. Renumber footers to stay consistent ───────────────────────────────────
n = len(prs.slides)
for idx, slide in enumerate(prs.slides):
    if idx == 0 or idx == n - 1:
        continue
    shapes = list(slide.shapes)
    footer = None
    for shp in reversed(shapes):
        if shp.has_text_frame and shp.text_frame.text.strip():
            footer = shp
            break
    if footer is not None:
        set_text_keep_format(footer, str(idx + 1))

prs.save(PPTX)
print("Saved.")
